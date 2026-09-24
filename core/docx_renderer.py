"""Word (.docx) 渲染器 —— 套用「模版.docx」设计系统。

设计 DNA（从模版 XML 精确提取）：
  主蓝 #2E75B6  → 分割线 / bullet 标记 / 段落下划线
  深蓝 #1F4E79  → 姓名 / 公司名 / 项目标题
  灰   #595959  → 副标题（部门/职位）/ 日期
  深灰 #404040  → 正文
  每条经历上方一条蓝色横线（卡片式分隔）；段落标题带蓝色下划线；
  bullet = 蓝色 • + 加粗小标题 + 关键短语加粗。

scale 等比缩放字号/标题/照片/段间距；dense=True 收紧行距/页边距以「一页 + 大字号」。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, replace
from io import BytesIO

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .parser import ResumeData

LATIN_FONT = "Arial"
BODY_CJK = "微软雅黑"
HEADING_CJK = "微软雅黑"

# —— 配色（模版）——
ACCENT = RGBColor(0x2E, 0x75, 0xB6)     # 主蓝
DARK_BLUE = RGBColor(0x1F, 0x4E, 0x79)   # 深蓝
GRAY = RGBColor(0x59, 0x59, 0x59)        # 副标题/日期
BODY_COLOR = RGBColor(0x40, 0x40, 0x40)  # 正文
LIGHT_GRAY = RGBColor(0x80, 0x80, 0x80)
ACCENT_HEX = "2E75B6"

CONTACT_LABELS = {
    "phone": "电话", "email": "邮箱", "github": "GitHub", "linkedin": "LinkedIn",
    "website": "网站", "location": "地址", "wechat": "微信", "other": "",
}


@dataclass
class _Cfg:
    name_size: float
    heading_size: float
    title_size: float
    contact_size: float
    date_size: float
    photo_w: float
    body_size: float
    sec_before: float
    sec_after: float
    entry_before: float
    bullet_after: float
    proj_before: float
    line_multiple: float
    m_top: float
    m_bottom: float
    m_lr: float


_BASE = _Cfg(
    name_size=20, heading_size=11.5, title_size=10, contact_size=9, date_size=9, photo_w=2.4,
    body_size=9.5, sec_before=7, sec_after=2, entry_before=3, bullet_after=1,
    proj_before=3, line_multiple=1.12,
    m_top=1.0, m_bottom=0.6, m_lr=1.2,
)


def _cfg_for(scale: float, dense: bool = False) -> _Cfg:
    b = _BASE
    c = replace(
        _BASE,
        name_size=b.name_size * scale, heading_size=b.heading_size * scale,
        title_size=b.title_size * scale, contact_size=b.contact_size * scale,
        date_size=b.date_size * scale, body_size=b.body_size * scale,
        photo_w=b.photo_w * scale,
        sec_before=b.sec_before * scale, sec_after=b.sec_after * scale,
        entry_before=b.entry_before * scale, bullet_after=b.bullet_after * scale,
        proj_before=b.proj_before * scale,
    )
    if dense:
        c = replace(
            c,
            line_multiple=1.0,
            m_top=0.5, m_bottom=0.35, m_lr=0.95,
            sec_before=c.sec_before * 0.55, sec_after=max(0.4, c.sec_after * 0.7),
            entry_before=c.entry_before * 0.55,
            bullet_after=max(0.2, c.bullet_after * 0.45),
            proj_before=c.proj_before * 0.6,
            photo_w=2.0,
        )
    return c


# —— 底层工具 ——
def _set_font(run, latin=LATIN_FONT, cjk=BODY_CJK, size=None, bold=None, color=None):
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rPr.get_or_add_rFonts().set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_inline_runs(paragraph, text, cjk=BODY_CJK, **font_kw):
    """解析 **加粗**；加粗段沿用基色（仅加粗不变色，对齐模版）。"""
    base_bold = font_kw.get("bold", False)
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not seg:
            continue
        run = paragraph.add_run(seg)
        font_kw["bold"] = base_bold or (i % 2 == 1)
        _set_font(run, cjk=cjk, **font_kw)


def _set_border(paragraph, edge: str, color: str = ACCENT_HEX, size: str = "8", space: str = "3"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), size)
    el.set(qn("w:space"), space)
    el.set(qn("w:color"), color)
    pbdr.append(el)


def _set_paragraph_shading(paragraph, fill: str):
    """段落底纹（整行背景色）——用于成果高亮蓝条。"""
    pPr = paragraph._p.get_or_add_pPr()
    shd = pPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        pPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def _set_cell_borderless(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)


def _is_project_header(detail: str) -> bool:
    s = detail.strip()
    return s.startswith("**") and s.endswith("**") and s.count("**") == 2


# —— 各部分渲染 ——
def _style_name(paragraph, resume, cfg):
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(resume.name)
    _set_font(run, cjk=HEADING_CJK, size=cfg.name_size, bold=True, color=DARK_BLUE)


def _style_contact(paragraph, resume, cfg):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    items = [(CONTACT_LABELS.get(k, ""), v) for k, v in resume.contact.items() if v]
    for i, (label, value) in enumerate(items):
        run = paragraph.add_run(f"{label} {value}".strip())
        _set_font(run, size=cfg.contact_size, color=GRAY)
        if i != len(items) - 1:
            sep = paragraph.add_run("  |  ")
            _set_font(sep, size=cfg.contact_size, color=LIGHT_GRAY)


def _render_header(doc, resume, photo_path, cfg, content_width_cm):
    if not photo_path:
        _style_name(doc.add_paragraph(), resume, cfg)
        _style_contact(doc.add_paragraph(), resume, cfg)
        return

    photo_col = cfg.photo_w + 0.3
    table = doc.add_table(rows=1, cols=2)
    table.allow_autofit = False
    left, right = table.cell(0, 0), table.cell(0, 1)
    left.width = Cm(content_width_cm - photo_col)
    right.width = Cm(photo_col)
    for cell in (left, right):
        _set_cell_borderless(cell)
    right.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    _style_name(left.paragraphs[0], resume, cfg)
    _style_contact(left.add_paragraph(), resume, cfg)

    pic_para = right.paragraphs[0]
    pic_para.alignment = 2  # RIGHT
    pic_para.paragraph_format.space_after = Pt(0)
    try:
        pic_para.add_run().add_picture(photo_path, width=Cm(cfg.photo_w))
    except Exception:
        pass


def _render_section_heading(doc, name, cfg):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(cfg.sec_before)
    pf.space_after = Pt(cfg.sec_after)
    icon = p.add_run("● ")
    _set_font(icon, size=cfg.heading_size - 3, bold=True, color=ACCENT)
    run = p.add_run(name)
    _set_font(run, cjk=HEADING_CJK, size=cfg.heading_size, bold=True, color=DARK_BLUE)
    _set_border(p, "bottom", color=ACCENT_HEX, size="6", space="2")


def _render_entry(doc, entry, cfg, content_width_cm, is_first: bool):
    has_details = bool(entry.details)

    # 技能型单行（无 details）：蓝色 • + 加粗标签
    if entry.title and not has_details:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(cfg.bullet_after)
        pf.line_spacing = cfg.line_multiple
        pf.left_indent = Cm(0.42)
        pf.first_line_indent = Cm(-0.42)
        pf.tab_stops.add_tab_stop(Cm(0.42), WD_TAB_ALIGNMENT.LEFT)
        marker = p.add_run("•\t")
        _set_font(marker, size=cfg.body_size, bold=True, color=ACCENT)
        _add_inline_runs(p, entry.title, size=cfg.body_size, color=BODY_COLOR)
        return

    # 条目抬头：公司(深蓝粗) + 部门/职位(灰) ←右对齐→ 日期(灰)
    if entry.title or entry.date:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(cfg.entry_before)
        pf.space_after = Pt(0.5)
        pf.tab_stops.add_tab_stop(Cm(content_width_cm), WD_TAB_ALIGNMENT.RIGHT)
        if not is_first:
            _set_border(p, "top", color=ACCENT_HEX, size="8", space="3")  # 卡片顶部蓝线
        if entry.title:
            _add_inline_runs(p, entry.title, cjk=HEADING_CJK, size=cfg.title_size, bold=True, color=DARK_BLUE)
        if entry.subtitle:
            sub = p.add_run("  " + entry.subtitle)
            _set_font(sub, size=cfg.body_size, color=GRAY)
        if entry.date:
            tab = p.add_run("\t" + entry.date)
            _set_font(tab, size=cfg.date_size, color=GRAY)

    for detail in entry.details:
        _render_bullet(doc, detail, cfg)


def _render_result_bar(doc, text, cfg):
    """成果高亮蓝条：浅蓝 #EBF3FB 底 + 左侧蓝竖条 #2E75B6 + 深蓝成果文字（对齐模版.docx）。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(1)
    pf.space_after = Pt(1)
    pf.line_spacing = cfg.line_multiple
    pf.left_indent = Cm(0.3)
    pf.right_indent = Cm(0.1)
    _set_paragraph_shading(p, "EBF3FB")
    _set_border(p, "left", color="2E75B6", size="18", space="4")
    _add_inline_runs(p, text, size=cfg.body_size, color=DARK_BLUE)
    return p


def _render_bullet(doc, detail, cfg):
    # 成果高亮蓝条：以 "> " 开头标记
    stripped = detail.strip()
    if stripped.startswith("> "):
        return _render_result_bar(doc, stripped[2:], cfg)
    # 项目小注：以 "~ " 开头 → 浅灰小字，紧贴项目标题
    if stripped.startswith("~ "):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(1)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.05)
        _add_inline_runs(p, stripped[2:], size=cfg.body_size - 1.5, color=LIGHT_GRAY)
        return
    # 项目子标题：【项目X】xxx → 深蓝加粗
    if _is_project_header(detail):
        text = detail.strip()[2:-2]
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(cfg.proj_before)
        pf.space_after = Pt(0.5)
        run = p.add_run(text)
        _set_font(run, cjk=HEADING_CJK, size=cfg.title_size, bold=True, color=DARK_BLUE)
        return

    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(cfg.bullet_after)
    pf.line_spacing = cfg.line_multiple
    pf.left_indent = Cm(0.42)
    pf.first_line_indent = Cm(-0.42)
    pf.tab_stops.add_tab_stop(Cm(0.42), WD_TAB_ALIGNMENT.LEFT)
    marker = p.add_run("•\t")
    _set_font(marker, size=cfg.body_size, bold=True, color=ACCENT)  # 蓝色 bullet
    _add_inline_runs(p, detail, size=cfg.body_size, color=BODY_COLOR)


def render_docx(resume_data: ResumeData, photo_path: str | None = None, scale: float = 1.0, dense: bool = False) -> bytes:
    cfg = _cfg_for(scale, dense=dense)
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(cfg.m_top)
    sec.bottom_margin = Cm(cfg.m_bottom)
    sec.left_margin = Cm(cfg.m_lr)
    sec.right_margin = Cm(cfg.m_lr)
    content_width_cm = 21.0 - 2 * cfg.m_lr

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(cfg.body_size)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_CJK)
    npf = normal.paragraph_format
    npf.space_before = Pt(0)
    npf.space_after = Pt(0)
    npf.line_spacing = cfg.line_multiple

    # 关闭段落对齐文档网格（默认模板 docGrid linePitch=360=18pt）：开启时段落高度会被
    # 量化到 18pt 整数倍，微调行距/段距会导致高度整体跳变（约 2 倍放大），无法精确排版。
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    normal.element.get_or_add_pPr().append(snap)

    _render_header(doc, resume_data, photo_path, cfg, content_width_cm)

    for section_name, entries in resume_data.sections.items():
        _render_section_heading(doc, section_name, cfg)
        for i, entry in enumerate(entries):
            _render_entry(doc, entry, cfg, content_width_cm, is_first=(i == 0))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
